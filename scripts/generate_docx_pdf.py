import os
import sys
import subprocess
import time

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(BASE_DIR, "docs", "screenshots")
DESKTOP_DIR = os.path.join(DOCS_DIR, "desktop")
MOBILE_DIR = os.path.join(DOCS_DIR, "mobile")

OUTPUT_DOCX = os.path.join(BASE_DIR, "DOKUMENTASI_PANDUAN_WASMAN.docx")
OUTPUT_PDF = os.path.join(BASE_DIR, "DOKUMENTASI_PANDUAN_WASMAN.pdf")

COLOR_PRIMARY = RGBColor(5, 150, 105)      # #059669
COLOR_PRIMARY_DARK = RGBColor(6, 78, 59)   # #064e3b
COLOR_SECONDARY = RGBColor(16, 185, 129)   # #10b981
COLOR_TEXT = RGBColor(31, 41, 55)          # #1f2937
COLOR_MUTED = RGBColor(107, 114, 128)      # #6b7280
HEX_PRIMARY = "059669"
HEX_PRIMARY_DARK = "064e3b"
HEX_BG_LIGHT = "ECFDF5"
HEX_GRAY_LIGHT = "F3F4F6"
HEX_BORDER = "D1D5DB"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def setup_document(doc):
    for s in doc.sections:
        s.page_width = Inches(8.27)    # A4
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Wasman — Platform Pengelolaan Sampah Berbasis Komunitas")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Dokumen Resmi Proyek Wasman v1.0 | Panduan Operasional 4 Tipe Pengguna")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED

    # Default Normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI Semibold"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_DARK
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI Semibold"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc, text, title="CATATAN OPERASIONAL", bg_hex=HEX_BG_LIGHT, border_hex=HEX_PRIMARY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.75)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    rt = p.add_run(f"📌 {title}\n")
    rt.font.bold = True
    rt.font.size = Pt(9.5)
    rt.font.color.rgb = COLOR_PRIMARY_DARK
    
    rb = p.add_run(text)
    rb.font.size = Pt(9.5)
    rb.font.color.rgb = COLOR_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_side_by_side_screenshots(doc, screenshot_name, title_caption):
    desk_path = os.path.join(DESKTOP_DIR, f"{screenshot_name}.png")
    mob_path = os.path.join(MOBILE_DIR, f"{screenshot_name}.png")
    
    if not os.path.exists(desk_path) or not os.path.exists(mob_path):
        p = doc.add_paragraph(f"[Screenshot {screenshot_name} tidak ditemukan]")
        return
        
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    r_cap = p_cap.add_run(f"Tangkapan Layar: {title_caption}")
    r_cap.font.bold = True
    r_cap.font.size = Pt(10)
    r_cap.font.color.rgb = COLOR_PRIMARY_DARK

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Desktop col ~ 4.4 in, Mobile col ~ 2.3 in
    table.columns[0].width = Inches(4.4)
    table.columns[1].width = Inches(2.35)
    
    # Header Row
    c_desk_hdr = table.cell(0, 0)
    c_mob_hdr = table.cell(0, 1)
    
    set_cell_background(c_desk_hdr, HEX_GRAY_LIGHT)
    set_cell_background(c_mob_hdr, HEX_GRAY_LIGHT)
    set_cell_margins(c_desk_hdr, top=60, bottom=60, left=100, right=100)
    set_cell_margins(c_mob_hdr, top=60, bottom=60, left=100, right=100)
    
    p0 = c_desk_hdr.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run("🖥️ Mode Website (Desktop 1280×850)")
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_TEXT
    
    p1 = c_mob_hdr.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run("📱 Mode Mobile (PWA 390×844)")
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_TEXT
    
    # Image Row
    c_desk_img = table.cell(1, 0)
    c_mob_img = table.cell(1, 1)
    
    set_cell_margins(c_desk_img, top=60, bottom=60, left=60, right=60)
    set_cell_margins(c_mob_img, top=60, bottom=60, left=60, right=60)
    
    p_img0 = c_desk_img.paragraphs[0]
    p_img0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img0.paragraph_format.space_after = Pt(0)
    p_img0.add_run().add_picture(desk_path, width=Inches(4.25))
    
    p_img1 = c_mob_img.paragraphs[0]
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_after = Pt(0)
    p_img1.add_run().add_picture(mob_path, width=Inches(2.15))
    
    # Border styling
    set_table_borders(table, color=HEX_BORDER, sz="4", val="single")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows, col_widths=None, header_bg=HEX_PRIMARY):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
            
    # Headers
    for i, h in enumerate(headers):
        c = table.cell(0, i)
        set_cell_background(c, header_bg)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row in enumerate(rows):
        bg = HEX_BG_LIGHT if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            c = table.cell(r_idx + 1, c_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=70, bottom=70, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_TEXT
            
    set_table_borders(table, color=HEX_BORDER, sz="4", val="single")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_cover_page(doc):
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(36)
    
    # Badge Box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.75)
    c = tbl.cell(0, 0)
    set_cell_background(c, HEX_PRIMARY)
    set_cell_margins(c, top=400, bottom=400, left=300, right=300)
    
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_sub = p.add_run("DOKUMEN SPESIFIKASI SISTEM & PANDUAN PENGGUNA\n")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(209, 250, 229)
    
    r_title = p.add_run("WASMAN\n")
    r_title.font.name = "Segoe UI Semibold"
    r_title.font.size = Pt(32)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    
    r_desc = p.add_run("Platform Pengelolaan Sampah Berbasis Komunitas\nMenuju Ekonomi Sirkular Berkelanjutan\n\n")
    r_desc.font.name = "Segoe UI"
    r_desc.font.size = Pt(14)
    r_desc.font.color.rgb = RGBColor(255, 255, 255)
    
    r_roles = p.add_run("Panduan Lengkap untuk 4 Tipe Pengguna:\n"
                       "1. Rumah Tangga (Household / Warga)\n"
                       "2. Komunitas (Bank Sampah / TPS 3R / Petugas)\n"
                       "3. Pabrik Pengolah (Offtaker / Recycler)\n"
                       "4. Administrator Sistem & DLH (Regulator)\n\n"
                       "Dilengkapi Komparasi Antarmuka Mode Website (Desktop) & Mode Mobile (PWA)")
    r_roles.font.name = "Segoe UI"
    r_roles.font.size = Pt(10.5)
    r_roles.font.color.rgb = RGBColor(209, 250, 229)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    
    # Metadata Footer Box
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta_p.add_run("Dokumen Resmi Arsitektur & Operasional Wasman v1.0\n"
                            "Disusun: September 2026 | Format Dokumen: Word DOCX & Adobe PDF")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = COLOR_MUTED
    
    doc.add_page_break()

def generate_document():
    print("Initializing document...")
    doc = docx.Document()
    setup_document(doc)
    
    # 1. Cover
    print("Generating Cover Page...")
    build_cover_page(doc)
    
    # 2. Ringkasan Eksekutif & Pengenalan Platform Wasman
    print("Writing Chapter 1: Pengenalan Platform...")
    add_h1(doc, "1. Ringkasan Eksekutif & Pengenalan Platform Wasman")
    
    doc.add_paragraph(
        "Wasman (Waste Management Platform) adalah sistem digital terintegrasi yang dirancang "
        "untuk mentransformasi tata kelola rantai pasok pengumpulan dan pengolahan sampah perkotaan "
        "menuju model Ekonomi Sirkular (Circular Economy). Platform ini menghubungkan penghasil sampah, "
        "pengangkut lokal, serta industri daur ulang dalam ekosistem berbasis data yang transparan, "
        "akuntabel, dan menguntungkan secara finansial."
    )
    
    add_h2(doc, "1.1 Urgensi Masalah & Solusi Inovatif")
    doc.add_paragraph(
        "Di berbagai wilayah perkotaan Indonesia, laju timbulan sampah belum diimbangi oleh tingkat pemilahan "
        "dari sumber yang memadai (< 15%). Hal ini mengakibatkan Tempat Pemrosesan Akhir (TPA) mengalami krisis "
        "kapasitas dan pencemaran lingkungan. Di sisi lain, industri daur ulang kekurangan bahan baku sekunder berkualitas "
        "dan harus mengimpor atau mengeluarkan biaya sortir tinggi akibat sampah tercampur dan kotor."
    )
    doc.add_paragraph(
        "Wasman menyelesaikan persoalan ini melalui 3 pendekatan strategis:"
    )
    doc.add_paragraph("• Insentif Langsung bagi Rumah Tangga: Memberikan reward berupa poin, uang tunai, atau potongan iuran kebersihan atas kepatuhan pemilahan sampah organik, anorganik, dan bernilai jual.")
    doc.add_paragraph("• Digitalisasi Operasional Komunitas: Memberdayakan Bank Sampah, TPS 3R, dan regu pengangkut lokal dengan sistem pengaturan jadwal, rute penjemputan cerdas, dan timbangan digital transparan.")
    doc.add_paragraph("• Kepastian Rantai Pasok Pabrik: Membuka saluran pasokan B2B langsung dari komunitas ke pabrik pengolah dengan standar mutu terukur (kemurnian > 95%) dan data asal material yang terlacak (traceable).")
    
    add_callout(doc, 
        "Prinsip Inti Ekonomi Sirkular Wasman:\n"
        "Sampah bukan limbah yang dibuang, melainkan material berharga (secondary raw material) "
        "yang dapat didaur ulang menjadi komoditas bernilai tambah jika dipilah sejak dari rumah tangga.",
        title="PRINSIP EKONOMI SIRKULAR"
    )
    
    add_h2(doc, "1.2 Matriks Peran 4 Stakeholders")
    headers_peran = ["Peran Pengguna", "Tipe Entitas", "Perangkat Utama", "Tanggung Jawab Kunci"]
    rows_peran = [
        ["1. Rumah Tangga", "Warga / Penghasil", "Smartphone (Mobile PWA)", "Memilah sampah dari sumber, konfirmasi jadwal, ajukan angkut khusus, lacak timbangan & kumpulkan reward."],
        ["2. Komunitas", "Bank Sampah / TPS 3R", "Mobile PWA & Web Desktop", "Agregator logistik, kelola armada & rute, timbang di lapangan, konsolidasi material, jual ke pabrik."],
        ["3. Pabrik", "Industri Daur Ulang", "Web Desktop (Utama) & Mobile", "Offtaker industri, rilis kuota B2B, uji mutu (QC), timbang jembatan timbang (bruto-tara-netto), bayar supplier."],
        ["4. Administrator", "Superadmin / DLH", "Web Desktop", "Regulator, tata kelola master data, harga acuan pasar, audit kepatuhan, verifikasi legalitas mitra, analitik kota."]
    ]
    add_styled_table(doc, headers_peran, rows_peran, [1.4, 1.4, 1.4, 2.55])
    
    add_h2(doc, "1.3 Halaman Utama Publik (Landing Page)")
    doc.add_paragraph(
        "Landing page publik Wasman menyajikan profil solusi, data dampak lingkungan terkini "
        "(tonase sampah terkelola, jumlah komunitas aktif, pabrik mitra), panduan cara kerja interaktif, "
        "serta tombol aksi pendaftaran bagi warga, komunitas, dan industri."
    )
    add_side_by_side_screenshots(doc, "landingpage", "Halaman Utama Publik (Landing Page)")
    
    # 3. Akses & Autentikasi
    print("Writing Chapter 2: Akses & Autentikasi...")
    add_h1(doc, "2. Akses & Autentikasi Sistem")
    doc.add_paragraph(
        "Wasman menerapkan arsitektur Single Sign-On (SSO) terpadu dengan pengalihan hak akses otomatis "
        "(Role-Based Access Control). Pengguna cukup memasukkan akun di satu pintu, dan sistem akan mengarahkan "
        "ke dasbor khusus sesuai perannya."
    )
    
    add_h2(doc, "2.1 Alur Masuk (Login Multi-Peran)")
    doc.add_paragraph(
        "1. Pengguna membuka URL sistem dan memasukkan email atau nomor WhatsApp terdaftar beserta kata sandi.\n"
        "2. Sistem memvalidasi status akun dan izin peran.\n"
        "3. Pengguna langsung diarahkan ke dasbor operasionalnya tanpa perlu memilih portal terpisah."
    )
    add_side_by_side_screenshots(doc, "auth-login", "Halaman Masuk (Login Multi-Peran)")
    
    add_h2(doc, "2.2 Alur Pendaftaran (Registrasi 3 Tipe Entitas)")
    doc.add_paragraph(
        "Formulir pendaftaran dirancang modular sesuai jenis entitas pendaftar:\n"
        "• Akun Warga Rumah Tangga: Mengisi data kepala keluarga, nomor WA, alamat domisili RT/RW, dan titik GPS.\n"
        "• Akun Komunitas / TPS 3R: Mengisi profil kelembagaan, penanggung jawab, zona wilayah binaan, dan armada.\n"
        "• Akun Pabrik: Mengisi nama badan usaha, izin lingkungan/NIB, alamat fasilitas, dan spesifikasi material olahan."
    )
    add_side_by_side_screenshots(doc, "auth-register", "Halaman Pendaftaran Akun Baru")
    
    # 4. Panduan Pengguna 1: Rumah Tangga
    print("Writing Chapter 3: Rumah Tangga...")
    add_h1(doc, "3. Panduan Pengguna 1: Rumah Tangga (Household / Warga)")
    doc.add_paragraph(
        "Rumah tangga merupakan garda terdepan sistem ekonomi sirkular. Keberhasilan daur ulang bertumpu pada "
        "kemauan warga untuk memisahkan sampah organik, anorganik bernilai, dan residu sejak di dapur rumah."
    )
    
    add_h2(doc, "3.1 Beranda (Dashboard) Warga")
    doc.add_paragraph(
        "Dashboard warga dirancang minimalis dan informatif untuk memantau performa pemilahan harian:\n"
        "• Kartu Statistik: Total frekuensi penyerahan sampah, total berat terpilah (kg), perolehan poin loyalitas, serta bintang reputasi pemilahan.\n"
        "• Pengingat Jadwal Terdekat: Menampilkan jadwal angkut rutin terdekat beserta tombol konfirmasi 'Saya Ada Sampah'.\n"
        "• Tips Pemilahan Harian: Panduan ringkas membersihkan sampah anorganik agar bernilai jual maksimal."
    )
    add_side_by_side_screenshots(doc, "household-dashboard", "Beranda Dashboard Rumah Tangga")
    
    add_h2(doc, "3.2 Memeriksa Jadwal Angkut & Konfirmasi 'Saya Ada Sampah'")
    doc.add_paragraph(
        "Pengangkutan reguler dijadwalkan oleh Komunitas pengelola wilayah setempat:\n"
        "1. Buka menu Jadwal Angkut pada bilah navigasi bawah (mobile) atau sidebar (desktop).\n"
        "2. Periksa jadwal operasional mingguan (misal: Senin - Organik, Rabu - Anorganik/Kardus, Jumat - Residu).\n"
        "3. Tekan tombol hijau 'Saya Ada Sampah' pada hari terkait jika telah memiliki sampah yang siap diangkut. "
        "Hal ini membantu petugas mengoptimalkan rute dan menghindari pemberhentian sia-sia."
    )
    add_side_by_side_screenshots(doc, "household-schedule", "Jadwal Angkut Reguler & Tombol Konfirmasi")
    
    add_h2(doc, "3.3 Pengajuan Pengangkutan Khusus / Bernilai")
    doc.add_paragraph(
        "Fitur ini digunakan saat warga memiliki material dalam jumlah besar atau kategori spesifik di luar jadwal reguler "
        "(kardus tebal, perabotan rusak, minyak jelantah, sampah elektronik/e-waste):\n"
        "1. Masuk ke menu 'Ajukan Permintaan' (ikon tambah).\n"
        "2. Pilih Kategori Sampah dan masukkan estimasi berat/volume.\n"
        "3. Pilih tanggal dan estimasi waktu penjemputan yang diinginkan.\n"
        "4. Unggah foto material sampah untuk mempermudah taksiran kapasitas kendaraan komunitas.\n"
        "5. Tulis catatan lokasi spesifik (misal: 'Kardus ditumpuk di dekat pagar depan rumah').\n"
        "6. Klik 'Kirim Permintaan' untuk diteruskan ke komunitas."
    )
    add_side_by_side_screenshots(doc, "household-collections-create", "Formulir Pengajuan Pengangkutan Khusus")
    
    add_h2(doc, "3.4 Riwayat & Pelacakan Status Pengangkutan")
    doc.add_paragraph(
        "Warga dapat memantau perjalanan tiket penjemputan secara real-time melalui tab Riwayat:\n"
        "• Menunggu Persetujuan: Permintaan sedang diverifikasi oleh operator komunitas.\n"
        "• Dijadwalkan: Penjemputan telah dialokasikan ke petugas dan armada.\n"
        "• Dalam Perjalanan: Petugas sedang bergerak menuju alamat rumah tangga.\n"
        "• Selesai: Penjemputan dan penimbangan telah berhasil disahkan."
    )
    add_side_by_side_screenshots(doc, "household-collections-index", "Daftar Riwayat & Pelacakan Status")
    
    add_h2(doc, "3.5 Detail Penjemputan, Timbangan Aktual, Insentif & Rating")
    doc.add_paragraph(
        "Setelah proses penjemputan selesai, warga dapat membuka tiket untuk memastikan transparansi:\n"
        "• Memeriksa berat aktual yang dicatat petugas melalui timbangan portabel digital.\n"
        "• Memeriksa pertambahan poin reward atau insentif uang tunai yang langsung masuk ke saldo akun.\n"
        "• Memberikan penilaian bintang (1-5) dan ulasan atas keramahan, ketepatan waktu, dan integritas timbangan petugas."
    )
    add_side_by_side_screenshots(doc, "household-collections-show", "Detail Tiket, Hasil Timbang & Rating Layanan")
    
    add_h2(doc, "3.6 Pusat Edukasi Pemilahan Sampah 3R")
    doc.add_paragraph(
        "Modul edukasi memberikan panduan praktis pengelolaan sampah:\n"
        "• Standar Kebersihan: Petunjuk membilas botol plastik, mengeringkan kardus, dan memisahkan tutup botol.\n"
        "• Kategori Material: Panduan membedakan sampah organik mudah terurai, daur ulang ekonomis, dan limbah B3 rumah tangga.\n"
        "• Artikel Lengkap: Penjelasan siklus hidup daur ulang dan dampaknya terhadap kelestarian lingkungan."
    )
    add_side_by_side_screenshots(doc, "household-education", "Katalog Panduan Edukasi 3R")
    add_side_by_side_screenshots(doc, "household-education-show", "Detail Artikel Edukasi Pemilahan")
    
    add_h2(doc, "3.7 Profil Akun & Notifikasi Real-Time")
    doc.add_paragraph(
        "Warga dapat mengelola informasi domisili, nomor kontak penjemputan, dan menerima notifikasi "
        "pengingat malam sebelum jadwal angkut, status kedatangan petugas lapangan, serta konfirmasi saldo reward."
    )
    add_side_by_side_screenshots(doc, "household-profile", "Pengaturan Profil & Alamat Jemput")
    add_side_by_side_screenshots(doc, "household-notifications", "Pusat Notifikasi Aktivitas Warga")
    
    # 5. Panduan Pengguna 2: Komunitas
    print("Writing Chapter 4: Komunitas...")
    add_h1(doc, "4. Panduan Pengguna 2: Komunitas (Community / Bank Sampah / TPS 3R)")
    doc.add_paragraph(
        "Komunitas memegang peran sentral sebagai agregator logistik lokal. Komunitas mengoordinasikan armada "
        "pengangkutan, menampung sampah terpilah di tempat penampungan sementara (TPS 3R), melakukan pengepakan material, "
        "dan memasoknya ke pabrik daur ulang."
    )
    
    add_h2(doc, "4.1 Dashboard Operasional Komunitas")
    doc.add_paragraph(
        "Dashboard komunitas memberikan ringkasan kendali operasional terpadu:\n"
        "• Metrik Kunci: Jumlah tugas lapangan hari ini, total RT binaan aktif, akumulasi tonase sampah bulan berjalan, dan skor kepuasan warga.\n"
        "• Panel 'Perlu Tindakan': Permintaan angkut khusus warga yang membutuhkan keputusan cepat (Setujui / Tolak).\n"
        "• Status Petugas Aktif: Memantau petugas lapangan yang sedang dalam perjalanan atau sedang istirahat.\n"
        "• Grafik Tren Mingguan: Fluktuasi volume sampah per hari untuk evaluasi rotasi armada."
    )
    add_side_by_side_screenshots(doc, "community-dashboard", "Dashboard Operasional Komunitas")
    
    add_h2(doc, "4.2 Pelaksanaan Tugas Lapangan Harian Petugas (Field Tasks)")
    doc.add_paragraph(
        "Halaman Tugas Lapangan dioptimalkan khusus untuk pengemudi armada motor roda tiga atau mobil bak di smartphone:\n"
        "1. Petugas membuka menu 'Tugas Hari Ini' di ponsel.\n"
        "2. Meninjau daftar titik jemput dengan status 'Pending'.\n"
        "3. Menekan tombol 'Mulai' saat bergerak menuju rumah warga. Status tiket otomatis berubah menjadi 'Dalam Perjalanan'.\n"
        "4. Tiba di lokasi, petugas menimbang sampah menggunakan timbangan gantung/duduk digital, mengambil foto bukti, dan memasukkan angka timbangan.\n"
        "5. Menekan tombol 'Selesai' untuk merampungkan tiket dan melanjutkan ke rumah berikutnya."
    )
    add_side_by_side_screenshots(doc, "community-tasks", "Lembar Tugas Lapangan Harian Petugas")
    
    add_h2(doc, "4.3 Manajemen Permintaan Masuk & Penugasan Petugas (Dispatching)")
    doc.add_paragraph(
        "Pengurus komunitas meninjau permintaan penjemputan warga dan menentukan alokasi sumber daya:\n"
        "1. Masuk ke menu 'Permintaan Angkut' dan pilih tiket yang berstatus 'Menunggu Persetujuan'.\n"
        "2. Klik 'Detail' untuk melihat foto sampah, estimasi berat, dan koordinat warga.\n"
        "3. Tentukan Petugas Lapangan yang bertugas dan pilih Armada Kendaraan yang sesuai.\n"
        "4. Klik 'Setujui & Jadwalkan'. Tiket otomatis masuk ke daftar tugas harian petugas terkait."
    )
    add_side_by_side_screenshots(doc, "community-collections-index", "Daftar Permintaan Penjemputan Masuk")
    add_side_by_side_screenshots(doc, "community-collections-show", "Detail Review Tiket & Penugasan Petugas")
    
    add_h2(doc, "4.4 Pengaturan Kalender & Pembuatan Jadwal Reguler")
    doc.add_paragraph(
        "Komunitas bertanggung jawab menyusun jadwal rutin mingguan agar kapasitas pengangkutan terkelola dengan baik:\n"
        "1. Buka menu 'Jadwal Angkut' ➔ Klik tombol 'Buat Jadwal Baru'.\n"
        "2. Pilih Wilayah Layanan sasaran (misal: RW 01 atau RW 02).\n"
        "3. Pilih Hari Pengangkutan dan rentang jam pelayanan.\n"
        "4. Pilih Kategori Sampah yang diangkut (Organik / Plastik & Kardus / Residu).\n"
        "5. Tentukan kuota maksimal rumah tangga yang dilayani pada sesi tersebut.\n"
        "6. Klik 'Simpan Jadwal'. Kalender langsung terdistribusi ke seluruh aplikasi warga di zona tersebut."
    )
    add_side_by_side_screenshots(doc, "community-schedules-index", "Kalender Jadwal Angkut Reguler")
    add_side_by_side_screenshots(doc, "community-schedules-create", "Formulir Pembuatan Jadwal Baru")
    
    add_h2(doc, "4.5 Manajemen Armada, Petugas & Wilayah Binaan")
    doc.add_paragraph(
        "Sistem menyediakan fasilitas tata kelola logistik internal:\n"
        "• Kelola Armada: Inventaris kendaraan (Tossa, Pickup, Truk Mini), kapasitas muatan (kg/m3), plat nomor, dan status servis.\n"
        "• Kelola Petugas: Data personil tim penjemput, nomor kontak, serta riwayat rating dan performa tugas.\n"
        "• Wilayah Layanan: Pemetaan zona RW/RT yang masuk dalam cakupan binaan operasional.\n"
        "• Database Rumah Tangga: Direktori warga terdaftar beserta rekam jejak konsistensi pemilahan."
    )
    add_side_by_side_screenshots(doc, "community-fleet", "Manajemen Armada Kendaraan")
    add_side_by_side_screenshots(doc, "community-officers", "Manajemen Petugas Lapangan")
    add_side_by_side_screenshots(doc, "community-service-areas", "Pengaturan Wilayah Layanan Binaan")
    add_side_by_side_screenshots(doc, "community-households", "Database Rumah Tangga Terdaftar")
    
    # 6. Panduan Pengguna 3: Pabrik Pengolah
    print("Writing Chapter 5: Pabrik Pengolah...")
    add_h1(doc, "5. Panduan Pengguna 3: Pabrik Pengolah / Recycler (Factory)")
    doc.add_paragraph(
        "Pabrik pengolah merupakan pembeli akhir (offtaker) yang menyerap material terpilah dalam skala tonase "
        "untuk dilebur, dicacah, atau diolah kembali menjadi produk baru. Pabrik menjamin keberlanjutan pasar daur ulang."
    )
    
    add_h2(doc, "5.1 Dashboard Industri Pabrik & Status Silo")
    doc.add_paragraph(
        "Dashboard pabrik dirancang untuk manajer pengadaan dan logistik industri:\n"
        "• Realisasi Pasokan Bulanan: Jumlah armada truk masuk, total tonase material diterima, jumlah komunitas pemasok, dan skor kemurnian rata-rata (Grade A/B/C).\n"
        "• Komposisi Material Diterima: Distribusi volume Botol PET Bening, Tutup HDPE, Kardus Karton, dan Gelas Minuman PP.\n"
        "• Leaderboard Pemasok Terbaik: Pemeringkatan komunitas mitra berdasarkan konsistensi tonase dan standar kemurnian sampah.\n"
        "• Status Jembatan Timbang & Keterisian Silo: Pemantauan sisa kapasitas tampung gudang dan antrean truk timbang."
    )
    add_side_by_side_screenshots(doc, "factory-dashboard", "Dashboard Industri & Kapasitas Silo Pabrik")
    
    add_h2(doc, "5.2 Profil Pabrik & Pengaturan Kebutuhan Material (B2B Demand)")
    doc.add_paragraph(
        "Pabrik dapat mengumumkan standar kebutuhan bahan baku daur ulang agar komunitas dapat menyesuaikan sortirannya:\n"
        "• Profil Perusahaan: Legalitas industri, izin lingkungan, dan kapasitas giling harian.\n"
        "• Spesifikasi Kebutuhan Bahan Baku: Menetapkan standar mutu kemurnian (Grade A > 95%), toleransi kontaminasi maksimum, dan harga beli per kg (misal: PET Bening Rp 3.000/kg, HDPE Rp 4.500/kg, Kardus Rp 1.800/kg).\n"
        "• Terbitkan Kebutuhan Borongan: Menerbitkan pesanan kuota bahan baku dengan target tanggal pengiriman."
    )
    add_side_by_side_screenshots(doc, "factory-profile", "Profil Pabrik & Katalog Spesifikasi Kebutuhan")
    
    add_h2(doc, "5.3 Riwayat Pasokan, Kontrol Mutu (QC) & Verifikasi Timbangan")
    doc.add_paragraph(
        "Prosedur penerimaan truk material dari komunitas di fasilitas pabrik:\n"
        "1. Verifikasi Surat Jalan Digital: Petugas gerbang memeriksa QR code pengiriman dari komunitas.\n"
        "2. Penimbangan Jembatan Timbang (Weighbridge): Menimbang truk beserta muatan (Bruto), truk membongkar material di silo, lalu truk kosong ditimbang kembali (Tara) untuk memperoleh Berat Bersih (Netto).\n"
        "3. Pengujian Laboratorium / Uji Sampel QC: Menghitung persentase kontaminasi pengotor atau kadar air.\n"
        "4. Pengesahan Invoice & Pembayaran: Nilai tagihan dihitung otomatis (Berat Bersih × Harga Satuan Mutu). Manajer pabrik menyetujui transfer pembayaran langsung ke rekening komunitas."
    )
    add_side_by_side_screenshots(doc, "factory-supply", "Riwayat Pasokan, Uji Mutu QC & Timbangan Pabrik")
    
    # 7. Panduan Pengguna 4: Administrator Sistem
    print("Writing Chapter 6: Administrator...")
    add_h1(doc, "6. Panduan Pengguna 4: Administrator Sistem & DLH (Admin)")
    doc.add_paragraph(
        "Administrator sistem atau Dinas Lingkungan Hidup (DLH) bertindak sebagai regulator ekosistem. "
        "Tugas utamanya mencakup penetapan master data, akreditasi legalitas mitra, pengawasan transparansi transaksi, "
        "serta analisis capaian target pengurangan sampah daerah."
    )
    
    add_h2(doc, "6.1 Executive Dashboard Kota & Reduksi Emisi Karbon")
    doc.add_paragraph(
        "Dashboard eksekutif merangkum metrik makro persampahan daerah secara real-time:\n"
        "• Agregat Makro: Total pengguna terdaftar, total rumah tangga aktif, jumlah bank sampah dan pabrik mitra, serta akumulasi timbulan sampah terkelola (ton).\n"
        "• Rasio Pemilahan Kota (Sorting Rate): Persentase sampah yang berhasil dicegah masuk ke TPA (landfill diversion rate).\n"
        "• Peringatan Sengketa (Dispute Alerts): Notifikasi dini terhadap tiket bermasalah (misal: selisih timbangan atau penolakan pabrik) untuk segera dimediasi.\n"
        "• Sebaran Geografis: Volume sampah terpilah per kecamatan untuk optimalisasi alokasi armada daerah."
    )
    add_side_by_side_screenshots(doc, "admin-dashboard", "Executive Dashboard Administrator Kota")
    
    add_h2(doc, "6.2 Tata Kelola Pengguna & Akreditasi Organisasi")
    doc.add_paragraph(
        "Untuk menjamin integritas dan keamanan ekosistem:\n"
        "• Kelola Pengguna: Audit status seluruh akun (Aktif, Menunggu Verifikasi, atau Ditangguhkan) lintas peran pengguna.\n"
        "• Verifikasi Organisasi Mitra: Memvalidasi berkas legalitas pendaftaran Bank Sampah baru, SK pembentukan TPS 3R, serta izin lingkungan AMDAL/NIB pabrik sebelum diizinkan bertransaksi."
    )
    add_side_by_side_screenshots(doc, "admin-users", "Tata Kelola Seluruh Akun Pengguna")
    add_side_by_side_screenshots(doc, "admin-organizations", "Verifikasi & Akreditasi Organisasi Mitra")
    
    add_h2(doc, "6.3 Master Data Kategori Sampah & Penetapan Harga Acuan Pasar")
    doc.add_paragraph(
        "Administrator menjaga standardisasi klasifikasi dan stabilitas harga material di pasaran:\n"
        "• Master Kategori Sampah: Mengatur taksonomi induk (Organik, Plastik, Kertas, Logam, Kaca, B3, Residu) dan prosedur keamanan penanganan.\n"
        "• Master Material & Harga Acuan: Menetapkan jenis material detail dan batas harga acuan (Floor Price & Ceiling Price) agar mencegah monopoli atau praktik perang harga yang merugikan komunitas pengumpul."
    )
    add_side_by_side_screenshots(doc, "admin-waste-categories", "Master Data Kategori Sampah")
    add_side_by_side_screenshots(doc, "admin-waste-materials", "Master Material & Penetapan Harga Acuan Pasar")
    
    # 8. SOP & Alur Bisnis Sirkular
    print("Writing Chapter 7: SOP & Alur Bisnis...")
    add_h1(doc, "7. Standard Operating Procedure (SOP) & Alur Bisnis Sirkular")
    doc.add_paragraph(
        "SOP ini mengikat seluruh pelaku ekosistem Wasman agar operasional pengumpulan dan pengolahan sampah "
        "berjalan transparan, tepat waktu, dan saling menguntungkan."
    )
    
    add_h2(doc, "7.1 SOP Pengangkutan Reguler Warga")
    doc.add_paragraph(
        "1. Komunitas merilis kalender jadwal rutin mingguan per wilayah RW/RT.\n"
        "2. Warga memilah sampah di rumah sesuai jadwal (Organik / Anorganik / Residu).\n"
        "3. Warga menekan tombol 'Saya Ada Sampah' di aplikasi sebelum batas jam konfirmasi (pukul 07.00 pagi).\n"
        "4. Petugas lapangan menerima peta rute rumah yang aktif dan memulai penjemputan.\n"
        "5. Petugas mengangkut sampah terpilah dan mencatat berat di sistem.\n"
        "6. Poin insentif warga diterbitkan otomatis setelah verifikasi selesai."
    )
    
    add_h2(doc, "7.2 SOP Penjemputan Sampah Khusus / Bernilai")
    doc.add_paragraph(
        "1. Warga membuat tiket permintaan dengan mengunggah foto material, estimasi berat, dan catatan alamat.\n"
        "2. Operator komunitas memeriksa kesiapan kapasitas armada dan memberikan persetujuan jadwal.\n"
        "3. Petugas lapangan mendatangi lokasi, melakukan penimbangan digital disaksikan warga, dan mengunggah foto bukti serah terima.\n"
        "4. Warga menerima notifikasi hasil timbang dan saldo reward bertambah seketika.\n"
        "5. Warga memberikan ulasan bintang 1-5 terhadap kepuasan layanan."
    )
    
    add_h2(doc, "7.3 SOP Pengiriman Pasokan Komunitas ke Pabrik Daur Ulang")
    doc.add_paragraph(
        "1. Komunitas mengonsolidasikan sampah terpilah di TPS 3R hingga mencapai kuota pengiriman truk.\n"
        "2. Komunitas membuat Tiket Pengiriman (Delivery Order) digital menuju pabrik tujuan.\n"
        "3. Truk tiba di pabrik, melewati jembatan timbang (Bruto), membongkar material, dan ditimbang kembali (Tara) untuk mendapatkan Netto murni.\n"
        "4. Tim QC pabrik menguji sampel kontaminasi. Jika kontaminasi < 5%, mutu disahkan sebagai Grade A.\n"
        "5. Pabrik menerbitkan invoice digital dan menyetujui transfer pembayaran ke rekening kas komunitas."
    )
    
    add_h2(doc, "7.4 Model Finansial & Transparansi Bagi Hasil")
    doc.add_paragraph(
        "Wasman menerapkan formula bagi hasil transparan yang menjamin seluruh rantai pasok mendapatkan margin yang adil. "
        "Berikut adalah simulasi pembagian nilai untuk 1 kg botol plastik PET bersih yang diserap pabrik seharga Rp 3.000 / kg:"
    )
    
    headers_fin = ["Alokasi Penerima", "Persentase", "Nilai (Rp/kg)", "Peruntukan & Fungsi Biaya"]
    rows_fin = [
        ["Insentif Warga (Rumah Tangga)", "50.0%", "Rp 1.500", "Reward pemilahan dari sumber (diberikan dalam bentuk saldo/poin belanja)."],
        ["Jasa Logistik & Sortir Komunitas", "36.7%", "Rp 1.100", "Biaya operasional armada, bahan bakar motor roda tiga, dan honor petugas lapangan."],
        ["Dana Kas Lingkungan RT/RW", "8.3%", "Rp 250", "Dana kebersihan warga, pembinaan lingkungan hidup, dan edukasi 3R lokal."],
        ["Biaya Platform Wasman", "5.0%", "Rp 150", "Pemeliharaan server cloud, layanan notifikasi WA, dan riset pengembangan sistem."],
        ["TOTAL NILAI TRANSAKSI PABRIK", "100.0%", "Rp 3.000", "Harga beli resmi pabrik untuk botol plastik PET kemurnian Grade A."]
    ]
    add_styled_table(doc, headers_fin, rows_fin, [1.8, 0.9, 1.0, 3.05])
    
    add_h2(doc, "7.5 Penanganan Sengketa & Pelanggaran (*Dispute Management*)")
    headers_sengketa = ["Jenis Masalah", "Tindakan Awal Sistem", "Tindakan Lanjutan / Sanksi"]
    rows_sengketa = [
        ["Sampah belum dipilah oleh warga", "Edukasi ramah dari petugas lapangan", "Petugas berhak menolak angkut / pengurangan poin insentif."],
        ["Warga membatalkan penjemputan berkali-kali", "Notifikasi peringatan sistem", "Pembatasan pengajuan angkut khusus selama 14 hari."],
        ["Selisih timbangan antara warga & petugas", "Verifikasi foto timbangan digital", "Penyesuaian saldo berimbang oleh admin komunitas."],
        ["Kontaminasi sampah tinggi di pabrik (> 10%)", "Penurunan grade mutu material (Grade B/C)", "Penyesuaian harga beli sesuai tabel mutu pabrik."],
        ["Pabrik menolak pasokan tanpa alasan valid", "Eskalasi ke dinas lingkungan (DLH)", "Pemeriksaan bukti timbang & sanksi penangguhan akun mitra industri."]
    ]
    add_styled_table(doc, headers_sengketa, rows_sengketa, [2.0, 2.2, 2.55])
    
    # 9. Penutup & Indeks Mockup
    print("Writing Chapter 8: Penutup & Indeks Mockup...")
    add_h1(doc, "8. Indeks Navigasi Mockup & Penutup")
    doc.add_paragraph(
        "Seluruh 31 halaman mockup prototipe Wasman dapat diakses secara langsung melalui berkas "
        "mockups/index.html. Struktur antarmuka telah dirancang responsif, modern, dan siap diimplementasikan "
        "ke dalam arsitektur produksi berbasis Laravel Modular Monolith, PWA, dan REST API."
    )
    add_side_by_side_screenshots(doc, "mockup-index", "Portal Navigasi Mockup Interaktif Wasman")
    
    add_callout(doc,
        "Rencana Tahapan Pengembangan Teknis (Roadmap MVP):\n"
        "• Fase 1 (MVP): Registrasi 4 pengguna, kalender jadwal reguler, modul edukasi 3R, dan pengajuan angkut khusus.\n"
        "• Fase 2 (Integrasi B2B): Portal pabrik, verifikasi jembatan timbang, penerbitan invoice otomatis, dan dompet digital insentif warga.\n"
        "• Fase 3 (Scale Up & IoT): Timbangan pintar terkoneksi IoT, pelacakan rute GPS real-time armada, dan kalkulator kredit jejak karbon.",
        title="ROADMAP PENGEMBANGAN TEKNIS",
        bg_hex=HEX_GRAY_LIGHT,
        border_hex=HEX_PRIMARY_DARK
    )
    
    print(f"Saving DOCX to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("DOCX successfully generated!")

def convert_to_pdf():
    print(f"Converting DOCX to PDF: {OUTPUT_PDF}...")
    ps_cmd = (
        f'$docx = "{OUTPUT_DOCX}"; '
        f'$pdf = "{OUTPUT_PDF}"; '
        f'$word = New-Object -ComObject Word.Application; '
        f'$word.Visible = $false; '
        f'try {{ '
        f'    $doc = $word.Documents.Open($docx); '
        f'    $doc.SaveAs([ref]$pdf, [ref]17); '
        f'    $doc.Close(); '
        f'    Write-Output "PDF_SUCCESS"; '
        f'}} finally {{ '
        f'    $word.Quit(); '
        f'    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null; '
        f'}}'
    )
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps_cmd], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if "PDF_SUCCESS" in res.stdout or os.path.exists(OUTPUT_PDF):
        sz = os.path.getsize(OUTPUT_PDF)
        print(f"PDF successfully generated: {OUTPUT_PDF} ({sz:,} bytes)")
    else:
        print(f"PDF conversion failed: {res.stderr} | {res.stdout}")

if __name__ == "__main__":
    t0 = time.time()
    generate_document()
    convert_to_pdf()
    print(f"All done in {time.time()-t0:.1f}s.")
